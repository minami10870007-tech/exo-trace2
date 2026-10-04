# -*- coding: utf-8 -*-
"""spec.py から詳細設計書（xlsx）を生成する。

使い方: python build.py [出力パス]
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import spec

FONT = "Arial"
F_BASE = Font(name=FONT, size=10)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=16, bold=True, color="1F4E5F")
F_SUB = Font(name=FONT, size=11, bold=True, color="1F4E5F")
F_LINK = Font(name=FONT, size=10, color="0563C1", underline="single")
FILL_HEAD = PatternFill("solid", fgColor="1F4E5F")
FILL_GROUP = PatternFill("solid", fgColor="D9E7EC")
FILL_KEY = PatternFill("solid", fgColor="FFF2CC")
FILL_INPUT = PatternFill("solid", fgColor="FFFF00")
THIN = Side(style="thin", color="A6A6A6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)

# シート名（目次・数式参照で共通利用）
S = {
    "cover": "表紙", "rev": "改訂履歴", "review": "レビュー記録", "overview": "システム概要",
    "flow": "業務フロー", "func": "機能一覧", "screen": "画面一覧", "sitem": "画面項目定義",
    "api": "API一覧", "proc": "処理詳細", "rule": "業務ルール", "state": "状態遷移",
    "table": "テーブル一覧", "column": "項目定義", "index": "インデックス・制約", "ddl": "DDL",
    "sql": "主要SQL", "code": "コード定義", "perm": "権限マトリクス", "msg": "メッセージ一覧",
    "report": "帳票一覧", "batch": "バッチ一覧", "nonfunc": "非機能要件", "issue": "課題一覧",
}
SHEET_DESC = {
    "rev": "版ごとの変更内容", "review": "自己レビュー（修正ループ）の指摘と対応",
    "overview": "目的・スコープ・前提・構成", "flow": "業務の流れと関連機能・データ",
    "func": "機能ID・優先度・関連画面", "screen": "画面ID・概要・操作", "sitem": "主要画面の入出力項目とチェック",
    "api": "REST API・許可ロール・処理詳細", "proc": "主要処理のステップ・トランザクション・エラー",
    "rule": "業務ルール・バリデーション", "state": "ロット/発注/受注/出荷/回収の状態遷移",
    "table": "テーブル一覧（ER概要）", "column": "全テーブルの項目定義", "index": "一意制約・CHECK・インデックス",
    "ddl": "PostgreSQL用 CREATE TABLE（項目定義から自動生成）", "sql": "トレース・回収抽出・整合性チェックSQL",
    "code": "コード値の定義", "perm": "機能×ロールの権限", "msg": "画面メッセージ",
    "report": "帳票・CSV出力", "batch": "定期バッチ", "nonfunc": "性能・セキュリティ・運用等", "issue": "未決事項・確認事項",
}


def setup(ws, headers, widths, rows, center_cols=(), title=None, note=None):
    """見出し付きの表を書き込む。title があれば1行目にタイトルを置き表は3行目から。"""
    start = 1
    if title:
        ws.cell(row=1, column=1, value=title).font = F_SUB
        if note:
            ws.cell(row=2, column=1, value=note).font = Font(name=FONT, size=9, italic=True, color="595959")
        start = 3
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=start, column=c, value=h)
        cell.font, cell.fill, cell.border, cell.alignment = F_HEAD, FILL_HEAD, BORDER, CENTER
    for r, row in enumerate(rows, start + 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font, cell.border = F_BASE, BORDER
            cell.alignment = CENTER if c in center_cols else WRAP
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    last = start + max(len(rows), 1)
    ws.freeze_panes = ws.cell(row=start + 1, column=1)
    ws.auto_filter.ref = f"A{start}:{get_column_letter(len(headers))}{last}"
    ws.sheet_view.zoomScale = 90
    return start


def back_link(ws, col):
    c = ws.cell(row=1, column=col, value="← 表紙へ")
    c.hyperlink = f"#'{S['cover']}'!A1"
    c.font = F_LINK


def build(out):
    wb = Workbook()
    ws = wb.active
    ws.title = S["cover"]

    # ---------------- 表紙 ----------------
    ws["A1"] = spec.SYSTEM_NAME
    ws["A1"].font = F_TITLE
    ws["A2"] = spec.DOC_TITLE
    ws["A2"].font = Font(name=FONT, size=14, bold=True)
    info = [("版数", spec.DOC_VERSION), ("作成日", spec.REVISIONS[0][1]),
            ("最終更新日", spec.REVISIONS[-1][1]), ("作成", "Claude（自動生成・自己レビュー済）"),
            ("生成元", "docs/design/tools/spec.py → build.py（設計変更は spec.py を編集し再生成）")]
    for i, (k, v) in enumerate(info, 4):
        ws.cell(row=i, column=1, value=k).font = Font(name=FONT, size=10, bold=True)
        ws.cell(row=i, column=2, value=v).font = F_BASE

    ws["A10"] = "設計規模サマリ（各シートから自動集計）"
    ws["A10"].font = F_SUB
    # 各シートの見出し行は title 付きのため3行目。データは4行目以降。
    summary = [
        ("機能数", f"=COUNTA('{S['func']}'!A4:A2000)"),
        ("  うち MUST", f"=COUNTIF('{S['func']}'!E:E,\"MUST\")"),
        ("画面数", f"=COUNTA('{S['screen']}'!A4:A2000)"),
        ("画面項目数", f"=COUNTA('{S['sitem']}'!A4:A2000)"),
        ("API数", f"=COUNTA('{S['api']}'!A4:A2000)"),
        ("テーブル数", f"=COUNTA('{S['table']}'!A4:A2000)"),
        ("カラム数（共通カラム含む）", f"=SUM('{S['table']}'!E:E)"),
        ("業務ルール数", f"=COUNTA('{S['rule']}'!A4:A2000)"),
        ("未対応の課題", f"=COUNTIF('{S['issue']}'!F:F,\"未対応\")"),
        ("レビュー指摘数（累計）", f"=COUNTA('{S['review']}'!A4:A2000)"),
        ("  うち 対応済", f"=COUNTIF('{S['review']}'!F:F,\"対応済\")"),
    ]
    for i, (k, f) in enumerate(summary, 11):
        ws.cell(row=i, column=1, value=k).font = F_BASE
        c = ws.cell(row=i, column=2, value=f)
        c.font, c.alignment = F_BASE, Alignment(horizontal="right")
        ws.cell(row=i, column=1).border = ws.cell(row=i, column=2).border = BORDER

    top = 11 + len(summary) + 1
    ws.cell(row=top, column=1, value="目次（クリックで移動）").font = F_SUB
    for c, h in enumerate(["No", "シート", "内容"], 1):
        cell = ws.cell(row=top + 1, column=c, value=h)
        cell.font, cell.fill, cell.border, cell.alignment = F_HEAD, FILL_HEAD, BORDER, CENTER
    for n, key in enumerate([k for k in S if k != "cover"], 1):
        r = top + 1 + n
        ws.cell(row=r, column=1, value=n).alignment = CENTER
        link = ws.cell(row=r, column=2, value=S[key])
        link.hyperlink = f"#'{S[key]}'!A1"
        link.font = F_LINK
        ws.cell(row=r, column=3, value=SHEET_DESC[key]).font = F_BASE
        for c in (1, 2, 3):
            ws.cell(row=r, column=c).border = BORDER
        ws.cell(row=r, column=1).font = F_BASE
    leg = top + len(S) + 2
    ws.cell(row=leg, column=1, value="凡例").font = F_SUB
    legends = [(FILL_HEAD, "表の見出し"), (FILL_GROUP, "グループ見出し（テーブル名・処理名等）"),
               (FILL_KEY, "トレーサビリティ上の重要項目（★）"), (FILL_INPUT, "課題一覧：ステータス／回答欄は記入可（プルダウン）")]
    for i, (fill, txt) in enumerate(legends, leg + 1):
        ws.cell(row=i, column=1).fill = fill
        ws.cell(row=i, column=1).border = BORDER
        ws.cell(row=i, column=2, value=txt).font = F_BASE
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 60

    def new(key):
        return wb.create_sheet(S[key])

    # ---------------- 改訂履歴 ----------------
    w = new("rev")
    setup(w, ["版", "日付", "区分", "内容", "作成/修正者"], [8, 12, 10, 90, 14], spec.REVISIONS,
          center_cols=(1, 2, 3), title="改訂履歴")
    back_link(w, 5)

    # ---------------- レビュー記録 ----------------
    w = new("review")
    setup(w, ["No", "回", "観点", "指摘内容", "対応内容", "状態", "対応版"], [7, 6, 18, 60, 60, 9, 8],
          [(f"R-{i:02d}",) + r for i, r in enumerate(spec.REVIEWS, 1)], center_cols=(1, 2, 6, 7),
          title="レビュー記録（自己修正ループ）",
          note="自動整合性チェック（tools/check.py）と、別エージェントによる設計レビューの指摘・対応を記録")
    back_link(w, 7)

    # ---------------- システム概要 ----------------
    w = new("overview")
    setup(w, ["項目", "内容"], [24, 110], spec.OVERVIEW, title="システム概要")
    for r in range(4, 4 + len(spec.OVERVIEW)):
        w.cell(row=r, column=1).font = Font(name=FONT, size=10, bold=True)
    back_link(w, 2)

    # ---------------- 業務フロー ----------------
    w = new("flow")
    setup(w, ["No", "業務", "ステップ", "担当", "内容", "関連機能ID", "生成/更新データ"],
          [7, 8, 16, 11, 70, 16, 40], spec.FLOWS, center_cols=(1, 2, 4), title="業務フロー",
          note="仕入 → 受入検品 → 保管 → 受注 → ロット引当・出荷 → 返品／回収 の流れ。全工程でロット番号を引き継ぐ")
    back_link(w, 7)

    # ---------------- 機能一覧 ----------------
    w = new("func")
    setup(w, ["機能ID", "大分類", "機能名", "概要", "優先度", "関連画面ID"], [9, 10, 22, 70, 9, 14],
          spec.FUNCTIONS, center_cols=(1, 2, 5, 6), title="機能一覧",
          note="優先度：MUST＝初期リリース必須／SHOULD＝初期リリース推奨／COULD＝将来")
    w.cell(row=3, column=1).value = "機能ID"
    back_link(w, 6)

    # ---------------- 画面一覧 ----------------
    w = new("screen")
    setup(w, ["画面ID", "画面名", "区分", "概要", "主な操作"], [9, 20, 9, 50, 50], spec.SCREENS,
          center_cols=(1, 3), title="画面一覧")
    back_link(w, 5)

    # ---------------- 画面項目定義 ----------------
    w = new("sitem")
    screen_name = {s[0]: s[1] for s in spec.SCREENS}
    rows = [(i[0], screen_name[i[0]]) + i[1:] for i in spec.SCREEN_ITEMS]
    setup(w, ["画面ID", "画面名", "No", "項目名", "種別", "入出力", "必須", "桁/形式", "対応テーブル.カラム",
              "初期値", "チェック・備考"], [9, 16, 5, 20, 12, 8, 9, 16, 34, 16, 50], rows,
          center_cols=(1, 3, 6, 7), title="画面項目定義",
          note="主要画面（入荷・検品・受注・出荷・追跡・回収）の項目。マスタ画面は項目定義シートのカラムをそのまま入力項目とする")
    back_link(w, 11)

    # ---------------- API一覧 ----------------
    w = new("api")
    setup(w, ["API ID", "メソッド", "パス", "概要", "関連機能ID", "許可ロール", "処理詳細"],
          [9, 12, 38, 38, 18, 30, 9], spec.APIS, center_cols=(1, 2, 7), title="API一覧",
          note="全APIはサーバー側で認証・権限（権限マトリクス）を検証する。更新系は楽観ロック（version）必須。エラーは {code, message_id, message} を返す")
    back_link(w, 7)

    # ---------------- 処理詳細 ----------------
    w = new("proc")
    rows, groups, prev = [], [], None
    for p in spec.PROCESSES:
        if p[0] != prev:
            groups.append(len(rows))
            prev = p[0]
        rows.append(p)
    setup(w, ["処理ID", "処理名", "ステップ", "内容", "エラー時メッセージ"], [8, 18, 8, 100, 14], rows,
          center_cols=(1, 3, 5), title="処理詳細",
          note="特記なき限り各処理は1トランザクション。在庫・ロットの更新は SELECT ... FOR UPDATE で行ロックし、エラー時は全体ロールバック")
    for g in groups:
        for c in range(1, 6):
            w.cell(row=4 + g, column=c).fill = FILL_GROUP
    back_link(w, 5)

    # ---------------- 業務ルール ----------------
    w = new("rule")
    setup(w, ["ルールID", "区分", "ルール", "関連機能ID"], [9, 14, 100, 18], spec.RULES,
          center_cols=(1, 2), title="業務ルール・バリデーション")
    back_link(w, 4)

    # ---------------- 状態遷移 ----------------
    w = new("state")
    setup(w, ["対象", "現状態", "イベント", "次状態", "条件・処理", "実行ロール"], [10, 30, 20, 24, 50, 12],
          spec.STATES, center_cols=(1, 6), title="状態遷移")
    back_link(w, 6)

    # ---------------- テーブル一覧 / 項目定義 ----------------
    wt = new("table")
    wc = new("column")
    col_rows, group_rows = [], []
    for t, (lname, kind, desc) in spec.TABLES.items():
        group_rows.append(len(col_rows))
        col_rows.append((t, f"【{lname}】{desc}", "", "", "", "", "", "", "", ""))
        cols = list(spec.COLUMNS[t])
        common = spec.COMMON_COLUMNS[:2] if t in spec.APPEND_ONLY_TABLES else spec.COMMON_COLUMNS
        for n, c in enumerate(cols + list(common), 1):
            col_rows.append((t, n) + tuple(c))
    headers = ["テーブル", "No", "物理名", "論理名", "型", "桁", "PK", "FK参照", "NOT NULL", "初期値", "備考"]
    col_rows = [r if len(r) == 11 else r + ("",) for r in col_rows]
    setup(wc, headers, [22, 5, 22, 22, 11, 7, 5, 24, 9, 12, 46], col_rows, center_cols=(2, 5, 6, 7, 9),
          title="項目定義", note=f"各テーブル末尾に共通カラム（created_at/by, updated_at/by, version）を付与。"
                               f"追記のみのテーブル（{', '.join(sorted(spec.APPEND_ONLY_TABLES))}）は created_* のみ")
    for g in group_rows:
        r = 4 + g
        wc.merge_cells(start_row=r, start_column=2, end_row=r, end_column=11)
        for c in range(1, 12):
            wc.cell(row=r, column=c).fill = FILL_GROUP
            wc.cell(row=r, column=c).font = Font(name=FONT, size=10, bold=True)
    for r in range(4, 4 + len(col_rows)):
        if "★" in str(wc.cell(row=r, column=11).value or ""):
            for c in range(1, 12):
                wc.cell(row=r, column=c).fill = FILL_KEY
    back_link(wc, 11)

    trows = []
    for n, (t, (lname, kind, desc)) in enumerate(spec.TABLES.items(), 1):
        trows.append((t, lname, kind, desc, None, None))
    setup(wt, ["物理名", "論理名", "区分", "概要", "カラム数", "主な参照先（FK）"], [24, 16, 14, 46, 9, 50],
          trows, center_cols=(3, 5), title="テーブル一覧",
          note="カラム数は項目定義シートから COUNTIF で自動集計（グループ見出し行を除く）")
    for i, t in enumerate(spec.TABLES, 4):
        wt.cell(row=i, column=5, value=f"=COUNTIF('{S['column']}'!A:A,A{i})-1")
        refs = sorted({c[5].split(".")[0] for c in spec.COLUMNS[t] if c[5]})
        wt.cell(row=i, column=6, value=", ".join(refs) if refs else "—")
        wt.cell(row=i, column=1).hyperlink = f"#'{S['column']}'!A{4 + group_rows[i - 4]}"
        wt.cell(row=i, column=1).font = F_LINK
    back_link(wt, 6)

    # ---------------- インデックス・制約 ----------------
    w = new("index")
    setup(w, ["テーブル", "種別", "名前", "対象カラム", "目的"], [22, 13, 28, 40, 60], spec.INDEXES,
          center_cols=(2,), title="インデックス・制約定義", note="PK・FK制約は項目定義に従い全て作成する（FK列には別途インデックスを作成）")
    back_link(w, 5)

    # ---------------- DDL ----------------
    w = new("ddl")
    ddl_lines = ddl()
    w.cell(row=1, column=1, value="DDL（PostgreSQL）— 項目定義・インデックス定義から自動生成").font = F_SUB
    back_link(w, 2)
    for i, line in enumerate(ddl_lines, 3):
        c = w.cell(row=i, column=1, value=line)
        c.font = Font(name="Courier New", size=9)
    w.column_dimensions["A"].width = 130

    # ---------------- 主要SQL ----------------
    w = new("sql")
    setup(w, ["SQL ID", "用途", "SQL"], [8, 26, 120], spec.SQLS, center_cols=(1,), title="主要SQL")
    for r in range(4, 4 + len(spec.SQLS)):
        w.cell(row=r, column=3).font = Font(name="Courier New", size=9)
        w.row_dimensions[r].height = 15 * (spec.SQLS[r - 4][2].count("\n") + 1)
    back_link(w, 3)

    # ---------------- コード定義 ----------------
    w = new("code")
    setup(w, ["コード種別", "コード", "名称", "備考"], [16, 16, 22, 50], spec.CODES, center_cols=(1,),
          title="コード定義", note="DBにはコード（英字）を格納し、画面・帳票では名称を表示する")
    back_link(w, 4)

    # ---------------- 権限マトリクス ----------------
    w = new("perm")
    fname = {f[0]: f[2] for f in spec.FUNCTIONS}
    rows = [(fid, fname[fid]) + tuple(spec.PERMISSIONS[fid]) for fid, *_ in spec.FUNCTIONS]
    headers = ["機能ID", "機能名"] + [f"{spec.ROLE_NAMES[r]}\n{r}" for r in spec.ROLES]
    setup(w, headers, [9, 24] + [11] * len(spec.ROLES), rows,
          center_cols=tuple(range(1, 3 + len(spec.ROLES))), title="権限マトリクス",
          note="C＝登録/更新/参照　A＝承認・判定/参照　R＝参照のみ　-＝利用不可")
    w.row_dimensions[3].height = 30
    for r in range(4, 4 + len(rows)):
        w.cell(row=r, column=2).alignment = WRAP
        for c in range(3, 3 + len(spec.ROLES)):
            v = w.cell(row=r, column=c).value
            if v in ("C", "A"):
                w.cell(row=r, column=c).fill = PatternFill("solid", fgColor="E2EFDA")
            elif v == "-":
                w.cell(row=r, column=c).font = Font(name=FONT, size=10, color="A6A6A6")
    back_link(w, 8)

    # ---------------- メッセージ・帳票・バッチ・非機能 ----------------
    w = new("msg")
    setup(w, ["メッセージID", "種別", "メッセージ", "発生条件"], [12, 8, 80, 30], spec.MESSAGES,
          center_cols=(1, 2), title="メッセージ一覧", note="{xxx} は実行時に値を埋め込む")
    back_link(w, 4)
    w = new("report")
    setup(w, ["帳票ID", "帳票名", "形式", "抽出条件", "主な出力項目", "関連機能ID"], [8, 20, 10, 24, 60, 14],
          spec.REPORTS, center_cols=(1, 3), title="帳票一覧")
    back_link(w, 6)
    w = new("batch")
    setup(w, ["バッチID", "バッチ名", "実行タイミング", "処理内容", "失敗時対応"], [8, 18, 14, 70, 36],
          spec.BATCHES, center_cols=(1, 3), title="バッチ一覧")
    back_link(w, 5)
    w = new("nonfunc")
    setup(w, ["区分", "項目", "要件"], [12, 20, 100], spec.NONFUNC, center_cols=(1,), title="非機能要件")
    back_link(w, 3)

    # ---------------- 課題一覧 ----------------
    w = new("issue")
    rows = [i + ("",) for i in spec.ISSUES]
    setup(w, ["No", "区分", "内容", "対応方針", "期限", "ステータス", "回答・決定事項（記入欄）"],
          [7, 9, 60, 40, 16, 10, 40], rows, center_cols=(1, 2, 6), title="課題一覧",
          note="黄色セル（ステータス・回答欄）に記入してください。例：ステータス=回答済、回答=「化粧品原料のみ取扱い」")
    dv = DataValidation(type="list", formula1='"未対応,確認中,回答済,クローズ"', allow_blank=False)
    w.add_data_validation(dv)
    for r in range(4, 4 + len(rows)):
        for c in (6, 7):
            w.cell(row=r, column=c).fill = FILL_INPUT
        dv.add(w.cell(row=r, column=6))
    w.cell(row=3, column=7).comment = Comment("決定内容を記入し、spec.py に反映して再生成する", "Claude")
    back_link(w, 7)

    # 印刷設定
    for sh in wb.worksheets:
        sh.page_setup.orientation = "landscape"
        sh.page_setup.fitToWidth = 1
        sh.page_setup.fitToHeight = 0
        sh.sheet_properties.pageSetUpPr.fitToPage = True
        sh.print_options.gridLines = False

    wb.save(out)


def _table_order():
    """FK依存順（参照先が先）に並べる。自己参照は無視。"""
    deps = {t: {c[5].split(".")[0] for c in list(spec.COLUMNS[t]) + spec.COMMON_COLUMNS if c[5]} - {t}
            for t in spec.TABLES}
    order, done = [], set()
    while len(order) < len(deps):
        ready = [t for t in spec.TABLES if t not in done and deps[t] <= done]
        if not ready:
            raise ValueError(f"FK循環: {set(deps) - done}")
        order += ready
        done |= set(ready)
    return order


def ddl():
    type_map = {"TIMESTAMP": "TIMESTAMPTZ"}
    out = ["-- 生成元: docs/design/tools/spec.py（手修正しないこと）", "-- 作成順は外部キーの依存順", ""]
    for t in _table_order():
        lname, _, desc = spec.TABLES[t]
        common = spec.COMMON_COLUMNS[:2] if t in spec.APPEND_ONLY_TABLES else spec.COMMON_COLUMNS
        cols = list(spec.COLUMNS[t]) + list(common)
        out.append(f"-- {lname}：{desc}")
        out.append(f"CREATE TABLE {t} (")
        defs = []
        for name, _, typ, size, pk, fk, nn, default, _ in cols:
            if pk == "Y":
                defs.append(f"  {name} BIGINT GENERATED ALWAYS AS IDENTITY PRIMARY KEY")
                continue
            d = f"  {name} {type_map.get(typ, typ)}" + (f"({size})" if size else "")
            if nn == "Y":
                d += " NOT NULL"
            if default == "現在日時":
                d += " DEFAULT now()"
            elif default in ("true", "false") or default.lstrip("-").isdigit():
                d += f" DEFAULT {default}"
            elif default and typ in ("VARCHAR", "CHAR"):
                d += f" DEFAULT '{default}'"
            if fk:
                ft, fc = fk.split(".")
                d += f" REFERENCES {ft}({fc})"
            defs.append(d)
        out += [d + ("," if i < len(defs) - 1 else "") for i, d in enumerate(defs)]
        out.append(");")
        out.append(f"COMMENT ON TABLE {t} IS '{lname}';")
        for name, lcol, *_ in cols:
            out.append(f"COMMENT ON COLUMN {t}.{name} IS '{lcol}';")
        out.append("")
    out.append("-- 一意制約・CHECK・インデックス")
    for t, kind, name, target, purpose in spec.INDEXES:
        if kind == "UNIQUE":
            out.append(f"ALTER TABLE {t} ADD CONSTRAINT {name} UNIQUE ({target});")
        elif kind == "UNIQUE(部分)":
            col, cond = target.split(" WHERE ")
            out.append(f"CREATE UNIQUE INDEX {name} ON {t} ({col}) WHERE {cond};")
        elif kind == "INDEX":
            out.append(f"CREATE INDEX {name} ON {t} ({target});")
        elif kind == "CHECK":
            expr = purpose.split("：", 1)[1] if "：" in purpose else purpose
            out.append(f"ALTER TABLE {t} ADD CONSTRAINT {name} CHECK ({expr});")
    return out


if __name__ == "__main__":
    out = Path(sys.argv[1] if len(sys.argv) > 1 else Path(__file__).resolve().parents[1] / "EXO-TRACE_詳細設計書.xlsx")
    build(out)
    schema = out.parent / "schema.sql"
    schema.write_text("\n".join(ddl()) + "\n", encoding="utf-8")
    print(out)
    print(schema)
